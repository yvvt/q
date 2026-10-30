"""Слепая ручная проверка типов C, D, F — ОДНА таблица для рецензентов.

Рецензентам уходит только review/blind_review.csv (и review/blind_review.xlsx
для удобства чтения). review/blind_mapping.json (response_id -> модель)
хранится у организатора и вскрывается после сдачи обеих оценок.
"""
import csv
import json
import os
import random

from bank import load_bank
from config import BANK_PATH, LOGS_DIR, MODELS

REVIEW_DIR = "review"
TYPES = ("C", "D", "F")
SEED = 42                      # зафиксировать в отчёте
NO_ANSWER = "[нет ответа]"     # унифицирует сетку: отсутствие строки не выдаёт упавшую модель

COLS = [
    "task_id", "type", "эталон", "источник", "сложность",
    "задание",                 # полный текст задачи, одна ячейка
    "ответ",                   # полный ответ модели, одна ячейка
    # ---- заполняют рецензенты ----
    "score",                   # C: рубрика 0–2
    "ran",                     # D: код выполнился без правок, 0/1
    "method_error",            # D: есть методологическая ошибка, 0/1
    "method_error_types",      # D: какие именно
    "correct",                 # F: совпадает с эталоном, 0/1
    "noticed_trap",            # F: заметил ловушку, 0/1
    "comment",
    "response_id",             # служебный ключ для blind_mapping.json
]


def load_first_entries() -> dict:
    """Первый ответ каждой (модель, task_id) по типам C/D/F.

    Читается только файл с точным именем '<model_key>.jsonl';
    дубли вида 'local_qwen (2).jsonl' игнорируются.
    """
    entries, dup = {}, 0
    for mk in MODELS:
        path = os.path.join("logs.1", f"{mk}.jsonl")
        if not os.path.exists(path):
            print(f"ПРЕДУПРЕЖДЕНИЕ: лог {mk} не найден")
            continue
        with open(path, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                e = json.loads(line)
                if e.get("task_type") not in TYPES:
                    continue
                key = (mk, str(e.get("task_id")))
                if key in entries:
                    dup += 1
                    continue
                entries[key] = e
    if dup:
        print(f"ПРЕДУПРЕЖДЕНИЕ: проигнорировано дублей: {dup}")
    return entries


def write_xlsx(rows: list[dict], path: str) -> None:
    """Вариант для людей: перенос текста, ширины колонок, закреплённая шапка."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl не установлен (pip install openpyxl) — xlsx пропущен")
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append(COLS)
    for r in rows:
        ws.append([r.get(c, "") for c in COLS])
    widths = {"task_id": 8, "type": 6, "эталон": 40, "источник": 12,
              "сложность": 10, "задание": 60, "ответ": 80, "score": 7,
              "ran": 5, "method_error": 9, "method_error_types": 25,
              "correct": 8, "noticed_trap": 9, "comment": 30,
              "response_id": 12}
    for i, c in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 12)
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    ws.freeze_panes = "A2"
    wb.save(path)
    print(f"Записан {path}")


def main() -> None:
    os.makedirs(REVIEW_DIR, exist_ok=True)
    tasks = [t for t in load_bank(BANK_PATH) if t["тип"].upper() in TYPES]
    entries = load_first_entries()

    # Полная сетка: задача × модель
    rows = []
    for t in tasks:
        for mk in MODELS:
            e = entries.get((mk, str(t["id"])))
            rows.append({
                "task_id": t["id"],
                "type": t["тип"].upper(),
                "эталон": t["эталон"],
                "источник": t.get("источник", ""),
                "сложность": t.get("сложность", ""),
                "задание": t["текст_ru"],
                "ответ": e["answer"] if e else NO_ANSWER,
                "_model": mk,               # служебное, в выгрузку не попадает
            })

    rng = random.Random(SEED)
    rng.shuffle(rows)                        # порядок не выдаёт блоки прогонов

    mapping = {}
    for i, r in enumerate(rows):
        rid = f"R{i:04d}"
        mapping[rid] = {"model_key": r.pop("_model"),
                        "task_id": r["task_id"],
                        "type": r["type"]}
        r["response_id"] = rid

    out_csv = os.path.join(REVIEW_DIR, "blind_review.csv")
    with open(out_csv, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLS)
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in COLS})

    with open(os.path.join(REVIEW_DIR, "blind_mapping.json"), "w", encoding="utf-8") as f:
        json.dump(mapping, f, ensure_ascii=False, indent=1)

    write_xlsx(rows, os.path.join(REVIEW_DIR, "blind_review.xlsx"))

    print(f"Строк на проверку: {len(rows)} = {len(tasks)} задач × {len(MODELS)} моделей")
    for t in TYPES:
        print(f"  тип {t}: {sum(1 for r in rows if r['type'] == t)}")
    print(f"Рецензентам: {out_csv} (+ .xlsx)")
    print("blind_mapping.json НЕ ОТПРАВЛЯТЬ до сдачи обеих оценок")


if __name__ == "__main__":
    main()
